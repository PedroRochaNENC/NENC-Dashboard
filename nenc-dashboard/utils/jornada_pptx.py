"""
Apresentação (PPTX) da Análise Geral da Jornada de Compra.

16:9 e fundo branco, com gráficos nativos do PowerPoint (botão direito →
Editar dados mexe nos números) e tabelas nativas. Mesmo recorte e mesmos
números da página e do PDF; as tabelas completas ficam no Excel. As imagens
do estudo (uma por loja, uma por marca) entram depois da gôndola e das
embalagens, inteiras e sem distorção.

Cor por função: a marca foco em violeta NENC e as demais em cinza (quem diz a
marca é o rótulo, não a cor); o funil usa a rampa de acento do claro para o
escuro, na ordem notou → examinou → retornou; atributos usam azul, laranja e
verde-água, os três primeiros tons da paleta de referência, que se separam
entre si também para quem tem daltonismo. Valores ausentes ficam sem barra
(nunca viram zero).
"""

import math
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import pandas as pd
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from utils import pptx_kit
from utils.jornada_export import export_filename, filters_text
from utils.jornada_format import fmt_number, fmt_pct
from utils.jornada_ingest import TASK_LABELS
from utils.jornada_metrics import decision_by_store, time_kpi
from utils.jornada_taxonomy import fold
from utils.pptx_kit import (  # noqa: F401 - reexportados para os testes
    ACCENT,
    BAND_FILL,
    BASE_BAR,
    BODY_BOTTOM,
    BODY_TOP,
    CATEGORICAL,
    CONTENT_WIDTH,
    FONT,
    FUNNEL,
    HEADER_FILL,
    INK,
    LEFT,
    MAX_TABLE_ROWS,
    MUTED,
    OTHER,
    RULE,
    SECONDARY,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    WHITE,
)

# Nomes antigos do kit (agora em utils/pptx_kit), usados aqui e nos testes.
_rgb = pptx_kit.rgb
_clean = pptx_kit.clean_text
_number = pptx_kit.chart_value
_dash = pptx_kit.dash
_luminance = pptx_kit.luminance
_label_on = pptx_kit.label_on
_style_run = pptx_kit.style_run
_add_runs = pptx_kit.add_runs
_textbox = pptx_kit.textbox
_bullet = pptx_kit.bullet
_write_items = pptx_kit.write_items
_markdown_items = pptx_kit.markdown_items
_Deck = pptx_kit.Deck
_cell = pptx_kit.table_cell
_table = pptx_kit.native_table
_kpis = pptx_kit.kpis
_chart_base = pptx_kit.chart_base
_axes = pptx_kit.bar_axes
_gap_width = pptx_kit.gap_width
_bar_chart = pptx_kit.bar_chart
_clustered_chart = pptx_kit.clustered_chart
_stacked_chart = pptx_kit.stacked_chart
_picture = pptx_kit.picture
_pictures = pptx_kit.picture_slides


def _pct(value) -> str:
    return _dash(fmt_pct(value))


def _num(value, digits: int = 1) -> str:
    return _dash(fmt_number(value, digits))


def _frame(value) -> pd.DataFrame:
    return value if isinstance(value, pd.DataFrame) else pd.DataFrame()


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def _cover(deck: _Deck, project: Dict, model: Dict, metrics: Dict, generated_at: str) -> None:
    slide = deck.prs.slides.add_slide(deck.layout)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SLIDE_HEIGHT)
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    focus = (model.get("meta") or {}).get("focus_brand") or ""
    _textbox(slide, Inches(0.9), Inches(2.1), Inches(11.6), Inches(0.45), "JORNADA DE COMPRA · ANÁLISE GERAL",
             size=14, bold=True, color=ACCENT)
    _textbox(slide, Inches(0.9), Inches(2.6), Inches(11.6), Inches(1.1), project.get("name") or "Projeto",
             size=44, bold=True)
    details = ["Marca foco: {}".format(focus)] if focus else []
    details.append("Gerado em {}".format(generated_at))
    if project.get("data_version") is not None:
        details.append("dados na versão {}".format(project["data_version"]))
    _textbox(slide, Inches(0.9), Inches(3.85), Inches(11.6), Inches(0.4), " · ".join(details), size=16,
             color=SECONDARY)
    _textbox(slide, Inches(0.9), Inches(4.35), Inches(11.6), Inches(0.6),
             "Recorte: {}".format(filters_text(metrics.get("filters"), model)), size=13, color=MUTED)


def _key_numbers(deck: _Deck, project: Dict, model: Dict, metrics: Dict) -> None:
    slide = deck.slide("Números-chave", "Só entram as gravações incluídas na análise.")
    sample = metrics.get("sample") or {}
    recordings = _frame(model.get("recordings"))
    uncoded = int((recordings["status"] == "nao_codificada").sum()) if not recordings.empty else 0
    time_label, time_value = time_kpi(metrics)
    _kpis(slide, BODY_TOP, [
        ("Participantes na análise", str(sample.get("participants", 0))),
        ("Gravações incluídas", str(sample.get("recordings", 0))),
        ("Não codificadas (fora)", str(uncoded)),
        (time_label, "{} s".format(fmt_number(time_value, 1)) if time_value == time_value else "—"),
    ])
    top = BODY_TOP + Inches(1.55)
    focus = (model.get("meta") or {}).get("focus_brand") or ""
    brand = _frame(metrics.get("brand"))
    if focus and not brand.empty and brand["is_focus"].any():
        _textbox(slide, LEFT, top, CONTENT_WIDTH, Inches(0.3),
                 "Share visual de {} por célula (média por participante)".format(focus), size=12, color=SECONDARY)
        _kpis(slide, top + Inches(0.38), [
            ("{} · n={}".format(row["cell"], int(row["n"])), _pct(row["share_mean"]))
            for _, row in brand[brand["is_focus"]].head(4).iterrows()
        ])
        top += Inches(1.95)
    questions = [line.strip(" -•\t") for line in str(project.get("questions") or "").splitlines()
                 if line.strip(" -•\t")]
    if questions:
        box = _textbox(slide, LEFT, top + Inches(0.1), CONTENT_WIDTH, BODY_BOTTOM - top - Inches(0.1))
        _write_items(box.text_frame, [("h", "Perguntas do estudo")] + [("b", q) for q in questions[:5]], 14)


def _findings(deck: _Deck, metrics: Dict) -> None:
    findings = metrics.get("findings") or []
    areas = (("Gôndola", ("gondola",)), ("Navegação e decisão", ("preco", "navegacao", "decisao")),
             ("Escolha", ("escolha",)), ("Embalagens", ("embalagem",)))
    items: List[Tuple[str, str]] = []
    for area, sections in areas:
        chosen = [f for f in findings if f.get("section") in sections][:6]
        if chosen:
            items.append(("h", area))
            items += [("b", "{}{}".format(f["text"], " (descritivo)" if f.get("strength") == "descritivo" else ""))
                      for f in chosen]
    deck.flow("Principais achados", items or [("p", "Sem achados para este recorte.")],
              subtitle="Frases geradas das métricas, com o n de cada célula.", size=15)


def _gondola(deck: _Deck, metrics: Dict, focus: str) -> None:
    brand = _frame(metrics.get("brand"))
    if brand.empty:
        return
    cells = list(brand.groupby("cell", sort=False))
    maximum = min(1.0, math.ceil((brand["share_mean"].max() or 0.1) * 10 + 0.5) / 10)
    gap = Inches(0.3)
    width = int((CONTENT_WIDTH - gap) / 2)
    height = Inches(2.55)
    for start in range(0, len(cells), 4):
        slide = deck.slide(
            "Share visual por marca",
            "Média, por participante, da fração da atenção às marcas que cada uma levou. "
            "{} em destaque; mesma escala em todos os quadros.".format(focus or "Marca foco"),
        )
        for index, (cell, rows) in enumerate(cells[start:start + 4]):
            ranked = rows.sort_values("share_mean", ascending=False)
            n = int(rows["n"].iloc[0])
            box = (LEFT + (index % 2) * (width + gap), BODY_TOP + (index // 2) * (height + Inches(0.12)),
                   width, height)
            _bar_chart(slide, box, ranked["brand"].tolist(), ranked["share_mean"].tolist(),
                       highlight=ranked["is_focus"].astype(bool).tolist(), maximum=maximum,
                       title="{} · n={}{}".format(cell, n, " (descritivo)" if n < 5 else ""))

    headers = ["Célula", "Marca", "n", "Share", "Notou", "Examinou", "Retornou", "1ª notada", "TTFF rel. (s)",
               "Índice"]
    widths = [3.1, 1.6, 0.6, 0.95, 0.95, 1.05, 1.05, 1.05, 1.1, 0.95]

    def table_rows(frame):
        return [[row["cell"], row["brand"], int(row["n"]), _pct(row["share_mean"]), _pct(row["reach"]),
                 _pct(row["examined"]), _pct(row["revisit"]), _pct(row["first_noticed"]),
                 _num(row["rel_ttff_median"]), _num(row["presence_index"], 2)] for _, row in frame.iterrows()]

    note = ("TTFF relativo = segundos depois da primeira marca vista. Índice = share ÷ fração da gôndola da "
            "marca (acima de 1, rende mais atenção que o espaço que ocupa).")
    focus_rows = brand[brand["is_focus"]] if focus else brand.iloc[0:0]
    if not focus_rows.empty:
        together = len(focus_rows) <= 6
        slide = deck.slide(
            "Funil de atenção — {}".format(focus),
            "Notou = olhou; examinou = ao menos o limiar de exame na marca; retornou = voltou a uma AOI dela.",
        )
        table_height = Inches(0.34 * (len(focus_rows) + 1)) if together else 0
        chart_height = BODY_BOTTOM - BODY_TOP - (table_height + Inches(0.55) if together else 0)
        _clustered_chart(
            slide, (LEFT, BODY_TOP, CONTENT_WIDTH, chart_height),
            focus_rows["cell"].tolist(),
            [("Notou", focus_rows["reach"].tolist()), ("Examinou", focus_rows["examined"].tolist()),
             ("Retornou", focus_rows["revisit"].tolist())],
            FUNNEL, maximum=1.0,
        )
        if together:
            top = BODY_TOP + chart_height + Inches(0.1)
            _textbox(slide, LEFT, top, CONTENT_WIDTH, Inches(0.3), note, size=10, color=MUTED)
            _table(slide, LEFT, top + Inches(0.38), headers, table_rows(focus_rows), widths)
        else:
            deck.table_slides("Primeira olhada e presença — {}".format(focus), headers,
                              table_rows(focus_rows), widths, subtitle=note)
    else:
        ordered = pd.concat([rows.sort_values("share_mean", ascending=False) for _, rows in cells])
        deck.table_slides("Funil, primeira olhada e presença", headers, table_rows(ordered), widths,
                          subtitle=note)

    sku = _frame(metrics.get("sku"))
    if not sku.empty:
        top = pd.concat([rows.sort_values("share_mean", ascending=False).head(5)
                         for _, rows in sku.groupby("cell", sort=False)])
        deck.table_slides(
            "Produtos mais vistos",
            ["Célula", "Produto", "Marca", "Share", "Alcance", "Tempo médio (s)"],
            [[row["cell"], row["product"], row["brand"], _pct(row["share_mean"]), _pct(row["reach"]),
              _num(row["dwell_mean_s"], 2)] for _, row in top.iterrows()],
            [3.2, 3.9, 1.7, 1.0, 1.0, 1.3],
            subtitle="Os cinco produtos com maior share em cada loja.",
        )


def _navigation(deck: _Deck, model: Dict, metrics: Dict) -> None:
    attributes = _frame(metrics.get("attributes"))
    if not attributes.empty:
        configured = (model.get("meta") or {}).get("dimensions") or {}
        dimensions = [d for d in configured if d in set(attributes["dimension"])] or sorted(
            attributes["dimension"].unique())
        gap = Inches(0.3)
        width = int((CONTENT_WIDTH - gap) / 2)
        for start in range(0, len(dimensions), 2):
            slide = deck.slide("Atributos dos produtos",
                               "Fração da atenção entre os produtos em que o atributo aparece no nome da AOI.")
            for index, dimension in enumerate(dimensions[start:start + 2]):
                rows = attributes[attributes["dimension"] == dimension]
                cells = list(dict.fromkeys(rows["cell"]))
                order = [v for v in configured.get(dimension, []) if v in set(rows["value"])]
                order += sorted(set(rows["value"]) - set(order))
                pivot = rows.pivot_table(index="cell", columns="value", values="share_mean", aggfunc="first")
                colors = [CATEGORICAL[i] if i < len(CATEGORICAL) else OTHER for i in range(len(order))]
                _stacked_chart(
                    slide, (LEFT + index * (width + gap), BODY_TOP, width, BODY_BOTTOM - BODY_TOP),
                    cells, [(value, [pivot.loc[c, value] if value in pivot.columns else None for c in cells])
                            for value in order],
                    colors, title=str(dimension).capitalize(),
                )

        if "presence" in attributes:
            deck.table_slides(
                "Atributos: atenção × presença na gôndola",
                ["Célula", "Atributo", "Valor", "n", "Share", "Presença", "Índice"],
                [[row["cell"], row["dimension"], row["value"], int(row["n_defined"]), _pct(row["share_mean"]),
                  _pct(row["presence"]), _num(row["presence_index"], 2)] for _, row in attributes.iterrows()],
                [3.4, 1.6, 1.6, 0.7, 1.3, 1.5, 1.2],
                subtitle="Índice = share ÷ presença do valor na gôndola: acima de 1, atenção além do espaço "
                         "que o valor ocupa; perto de 1, proporcional.",
            )

    price = _frame(metrics.get("price"))
    if not price.empty:
        deck.table_slides(
            "Etiquetas de preço",
            ["Célula", "Produto", "Viu o preço", "Tempo médio (s)", "Preço ÷ (preço + produto)"],
            [[row["cell"], row["product"], _pct(row["reach"]), _num(row["dwell_mean_s"], 2),
              _pct(row["price_fraction"])] for _, row in price.iterrows()],
            [3.3, 3.8, 1.5, 1.6, 2.0],
            subtitle="Atenção às etiquetas de preço, onde estão mapeadas.",
        )

    decision = _frame(metrics.get("decision"))
    if not decision.empty:
        slide = deck.slide(
            "Tempo até a decisão",
            "Mediana por tarefa e loja, em segundos: o tempo de compra do registro de campo e, onde ele não "
            "existe, o Tempo da planilha.",
        )
        stores = decision_by_store(decision)
        if not stores.empty:
            _bar_chart(slide, (LEFT, BODY_TOP, Inches(5.6), Inches(3.4)), stores["label"].tolist(),
                       stores["median_s"].tolist(), number_format='0.0" s"', title="Mediana por tarefa e loja")
        table = decision[decision["group_type"] == "loja"]
        _table(slide, LEFT + Inches(6.0), BODY_TOP, ["Tarefa", "Fonte", "Loja", "n", "Mediana (s)"],
               [[row["task_label"], row["source_label"], row["group"], int(row["n"]), _num(row["median_s"])]
                for _, row in table.head(MAX_TABLE_ROWS).iterrows()],
               [1.5, 1.85, 1.15, 0.45, 1.15])


def _choice(deck: _Deck, metrics: Dict) -> None:
    choices = _frame(metrics.get("choices"))
    choice = _frame(metrics.get("choice"))
    if choices.empty:
        return
    subtitle = ("Registro de campo, normalizado pelas marcas do projeto. Livre = compra observada; estimulada = "
                "escolha pedida. Quem escolheu duas marcas conta nas duas.")
    for _, rows in choice.groupby("task", sort=False):
        total = rows[rows["group_type"] == "total"].sort_values("share", ascending=False)
        if total.empty:
            continue
        task_label = rows["task_label"].iloc[0]
        n = int(total["n"].iloc[0])
        slide = deck.slide("Escolha — {}".format(task_label), subtitle)
        _bar_chart(slide, (LEFT, BODY_TOP, Inches(5.6), BODY_BOTTOM - BODY_TOP), total["brand"].tolist(),
                   total["share"].tolist(), highlight=total["is_focus"].astype(bool).tolist(), maximum=1.0,
                   title="Todas as lojas · n={}{}".format(n, " (descritivo)" if n < 5 else ""))
        stores = rows[(rows["group_type"] == "loja") & (rows["chose_n"] > 0)].head(MAX_TABLE_ROWS)
        _textbox(slide, LEFT + Inches(6.0), BODY_TOP, Inches(6.0), Inches(0.35), "Por loja", size=13, bold=True)
        _table(slide, LEFT + Inches(6.0), BODY_TOP + Inches(0.45), ["Loja", "Marca", "Escolheram", "Fração"],
               [[row["group"], row["brand"], "{}/{}".format(int(row["chose_n"]), int(row["n"])), _pct(row["share"])]
                for _, row in stores.iterrows()],
               [2.1, 1.7, 1.2, 1.0])
        # Canal em texto, abaixo da tabela: na mesma tabela as lojas empurravam os canais para fora.
        channels = rows[rows["group_type"] == "canal"]
        lines = ["{} — {}".format(channel, " · ".join(
            "{} {}/{}".format(row["brand"], int(row["chose_n"]), int(row["n"])) for _, row in part.iterrows()))
            for channel, part in channels.groupby("group", sort=False)]
        if lines:
            top = BODY_TOP + Inches(0.45) + Inches(0.34 * (len(stores) + 1)) + Inches(0.2)
            box = _textbox(slide, LEFT + Inches(6.0), top, Inches(6.0), BODY_BOTTOM - top)
            _write_items(box.text_frame, [("h", "Por canal")] + [("b", line) for line in lines], 12)
    attention = _frame(metrics.get("attention_choice"))
    consideration = _frame(metrics.get("consideration"))
    if attention.empty and consideration.empty:
        return
    slide = deck.slide("Da atenção à escolha",
                       "Quem tem escolha e olhar codificado na mesma gravação: a marca escolhida foi notada, "
                       "examinada, a primeira e a mais vista?")
    top = BODY_TOP
    if not attention.empty:
        n = len(attention)
        _kpis(slide, top, [
            ("Notaram a escolhida", "{}/{}".format(int(attention["looked"].sum()), n)),
            ("Examinaram", "{}/{}".format(int(attention["examined"].sum()), n)),
            ("Foi a 1ª notada", "{}/{}".format(int(attention["first_noticed"].sum()), n)),
            ("Foi a mais vista", "{}/{}".format(int(attention["top_share"].sum()), n)),
        ])
        top += Inches(1.6)
    if not consideration.empty:
        _textbox(slide, LEFT, top, CONTENT_WIDTH, Inches(0.35), "Conjunto considerado e embalagens citadas",
                 size=13, bold=True)
        _table(slide, LEFT, top + Inches(0.45),
               ["Tarefa", "Loja", "n", "Marcas consideradas (média)", "Só uma marca", "Embalagens citadas"],
               [[row["task_label"], row["store_label"], int(row["n"]), _num(row["considered_mean"]),
                 int(row["single_brand_n"]), _dash(row["packs"])]
                for _, row in consideration.head(8).iterrows()],
               [2.0, 1.8, 0.6, 2.2, 1.4, 4.1])


def _packaging(deck: _Deck, metrics: Dict, focus: str) -> None:
    packaging = metrics.get("packaging") or {}
    elements = _frame(packaging.get("elements"))
    if elements.empty:
        return
    brands = _frame(packaging.get("brands"))
    coverage = _frame(packaging.get("coverage"))
    group = "Todos" if (elements["group"] == "Todos").any() else elements["group"].iloc[0]
    rows = elements[elements["group"] == group]
    subtitle = "Agregado por perfil, sem variação entre participantes · base: {}, n={}".format(
        rows["profile"].iloc[0], int(rows["n_group"].max()))
    cov = coverage[coverage["group"] == group] if not coverage.empty else coverage
    if not cov.empty and cov["aoi_coverage"].notna().any():
        subtitle += " · elementos somam {} do tempo gravado".format(_pct(cov["aoi_coverage"].iloc[0]))
    slide = deck.slide("Embalagens", subtitle)
    focus_rows = rows[rows["brand"].map(fold) == fold(focus)] if focus else rows.iloc[0:0]
    if not focus_rows.empty:
        ranked = focus_rows.sort_values("element_share", ascending=False)
        _bar_chart(slide, (LEFT, BODY_TOP, Inches(6.2), BODY_BOTTOM - BODY_TOP), ranked["element_label"].tolist(),
                   ranked["element_share"].tolist(), title="Onde o olhar cai na embalagem de {}".format(focus))
    summary = []
    for brand_name, brand_rows in rows.groupby("brand", sort=False):
        top = brand_rows.sort_values("element_share", ascending=False).iloc[0]
        logo = brands[(brands["group"] == group) & (brands["brand"] == brand_name)] if not brands.empty else brands
        summary.append([brand_name, top["element_label"], _pct(top["element_share"]),
                        _pct(logo["logo_reach"].iloc[0]) if not logo.empty else "—"])
    _textbox(slide, LEFT + Inches(6.6), BODY_TOP, Inches(5.5), Inches(0.35), "Elemento que mais atrai",
             size=13, bold=True)
    _table(slide, LEFT + Inches(6.6), BODY_TOP + Inches(0.45), ["Marca", "Elemento mais visto", "Share", "Logo visto"],
           summary[:MAX_TABLE_ROWS], [1.45, 1.95, 0.9, 1.2])

    if focus and not brands.empty:
        logos = brands[brands["brand"].map(fold) == fold(focus)]
        if len(logos) > 1:
            slide = deck.slide("Logo de {} visto, por perfil".format(focus),
                               "Fração do grupo que olhou o elemento de marca da embalagem.")
            _bar_chart(slide, (LEFT, BODY_TOP, CONTENT_WIDTH, Inches(4.2)),
                       ["{} (n={})".format(row["profile"], int(row["n_group"])) for _, row in logos.iterrows()],
                       logos["logo_reach"].tolist(), highlight=(logos["group"] == "Todos").tolist(), maximum=1.0)


def _channel(deck: _Deck, metrics: Dict) -> None:
    confounds = metrics.get("confounds") or []
    comparisons = _frame(metrics.get("comparisons"))
    if not confounds and comparisons.empty:
        return
    slide = deck.slide("Canal e perfil", "Com menos de 5 participantes por grupo a comparação é só descritiva.")
    top = BODY_TOP
    if confounds:
        box = _textbox(slide, LEFT, top, CONTENT_WIDTH, Inches(0.4 + 0.45 * len(confounds)))
        _write_items(box.text_frame, [
            ("b", "**{} e {} andam juntos nesta amostra** ({}): uma diferença entre um pode ser do outro, e a "
                  "comparação não isola o efeito.".format(first, second.lower(), confound["mapping"]))
            for confound in confounds for first, second in [confound["labels"]]
        ], 14)
        top += Inches(0.55 + 0.45 * len(confounds))
    if not comparisons.empty:
        _table(slide, LEFT, top,
               ["Tarefa", "Por", "Grupo A", "Grupo B", "Média A", "Média B", "δ de Cliff", "p", "Método"],
               [[TASK_LABELS.get(row["task"], row["task"]), row["by"],
                 "{} (n={})".format(row["group_a"], row["n_a"]), "{} (n={})".format(row["group_b"], row["n_b"]),
                 _pct(row["mean_a"]), _pct(row["mean_b"]), _num(row["cliffs_delta"], 2), _num(row["p_value"], 3),
                 row["method"]] for _, row in comparisons.head(MAX_TABLE_ROWS - 2).iterrows()],
               [1.6, 0.7, 2.1, 2.1, 0.95, 0.95, 1.05, 0.7, 1.95])


def _sample(deck: _Deck, model: Dict, metrics: Dict, quality: Optional[Dict]) -> None:
    subtitle = ""
    summary = _frame((quality or {}).get("summary"))
    if not summary.empty:
        tally = summary["quality"].value_counts()
        subtitle = "Qualidade das gravações: {} OK, {} com atenção, {} com problema.".format(
            int(tally.get("pass", 0)), int(tally.get("warn", 0)), int(tally.get("fail", 0)))
    slide = deck.slide("Amostra e qualidade", subtitle)
    by_cell = (metrics.get("sample") or {}).get("by_cell") or []
    if by_cell:
        _table(slide, LEFT, BODY_TOP, ["Célula", "Participantes"],
               [[item["cell"], item["n"]] for item in by_cell[:MAX_TABLE_ROWS]], [3.6, 1.5])
    recordings = _frame(model.get("recordings"))
    if not recordings.empty:
        titles = {"incluida": "Incluída", "nao_codificada": "Não codificada", "excluida": "Excluída",
                  "agregado": "Só agregado", "sem_aoi": "Sem AOI"}
        counts = recordings.groupby(["task_label", "status"]).size().unstack(fill_value=0)
        statuses = [s for s in titles if s in counts.columns]
        _table(slide, LEFT + Inches(5.5), BODY_TOP, ["Tarefa"] + [titles[s] for s in statuses],
               [[task] + [int(counts.loc[task, s]) for s in statuses] for task in counts.index],
               [1.9] + [4.7 / max(1, len(statuses))] * len(statuses))
    issues = [i for i in model.get("issues") or [] if i.get("level") in ("warn", "error")]
    if issues:
        box = _textbox(slide, LEFT, Inches(4.6), CONTENT_WIDTH, Inches(2.2))
        items = [("h", "Avisos dos dados")] + [("b", i["message"]) for i in issues[:5]]
        if len(issues) > 5:
            items.append(("p", "Mais {} aviso(s) no Excel (aba Avisos).".format(len(issues) - 5)))
        _write_items(box.text_frame, items, 12)


def _analysis(deck: _Deck, analysis: Dict, data_version) -> None:
    modes = {"rapida": "modo rápido", "aprofundada": "modo aprofundado"}
    details = [analysis.get("model") or "", modes.get(analysis.get("mode"), analysis.get("mode") or ""),
               "gerada em {}".format(analysis["created_at"]) if analysis.get("created_at") else ""]
    subtitle = " · ".join(d for d in details if d)
    if analysis.get("data_version") is not None and analysis.get("data_version") != data_version:
        subtitle += " · os dados mudaram depois desta análise"
    items = _markdown_items(analysis.get("analysis_text") or "")
    citations = analysis.get("citations") or []
    if citations:
        items.append(("h", "Referências"))
        for citation in citations[:8]:
            data = citation if isinstance(citation, dict) else {"quote": str(citation)}
            quote = str(data.get("quote") or "").strip().replace("**", "")
            if len(quote) > 200:
                quote = quote[:200].rstrip() + "…"
            name = str(data.get("filename") or "Documento").replace("**", "")
            items.append(("b", "**{}**{}".format(name, " — “{}”".format(quote) if quote else "")))
    deck.flow("Análise de IA", items, subtitle=subtitle, size=14)


def build_pptx(
    project: Dict,
    model: Dict,
    metrics: Dict,
    *,
    analysis: Optional[Dict] = None,
    quality: Optional[Dict] = None,
    generated_at: Optional[str] = None,
    images: Optional[List[Dict]] = None,
) -> Tuple[bytes, str]:
    """Apresentação da Análise Geral no recorte da página; devolve os bytes e o nome.

    `images` vem de `jornada_gallery.report_images`, com `content` (bytes já reduzidos).
    """

    generated_at = generated_at or datetime.now().strftime("%d/%m/%Y %H:%M")
    focus = (model.get("meta") or {}).get("focus_brand") or ""
    deck = _Deck("{} · Jornada de Compra · {}".format(project.get("name") or "Projeto", generated_at))
    _cover(deck, project, model, metrics, generated_at)
    _key_numbers(deck, project, model, metrics)
    _findings(deck, metrics)
    pictures = images or []
    _gondola(deck, metrics, focus)
    _pictures(deck, [item for item in pictures if item.get("group") == "loja"], "A gôndola de cada loja",
              "Heatmap da loja quando existe; senão, a foto da gôndola.")
    _navigation(deck, model, metrics)
    _choice(deck, metrics)
    _packaging(deck, metrics, focus)
    _pictures(deck, [item for item in pictures if item.get("group") == "marca"], "As embalagens",
              "A versão editada de cada marca quando existe; senão, a frente.")
    _channel(deck, metrics)
    _sample(deck, model, metrics, quality)
    if analysis and analysis.get("analysis_text"):
        _analysis(deck, analysis, project.get("data_version"))
    deck.flow("Limitações", [("b", note) for note in metrics.get("limitations") or ["Sem limitações registradas."]],
              size=15)
    return deck.save(), export_filename(project, "pptx", "jornada_apresentacao")
