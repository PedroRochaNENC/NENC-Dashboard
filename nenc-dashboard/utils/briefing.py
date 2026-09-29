"""
Leitura do documento de briefing de um projeto.

O NencBoost e a Jornada de Compra aceitam o mesmo tipo de material para dar
contexto a IA; antes cada tela tinha o proprio extrator, e so um deles lia
PowerPoint. As funcoes sao puras (bytes entram, texto sai) para servirem as
duas telas e os testes sem Streamlit.
"""

import io
import xml.etree.ElementTree as ET
import zipfile
from typing import Tuple

BRIEFING_EXTENSIONS = ("txt", "md", "csv", "json", "docx", "pptx")
BRIEFING_MAX_CHARS = 20000

_WORD_NAMESPACE = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def decode_text_bytes(data: bytes) -> str:
    """Texto de um arquivo cuja codificacao ninguem informou."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_docx_text(data: bytes) -> str:
    """Paragrafos de um .docx, sem depender do python-docx."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml_bytes = archive.read("word/document.xml")

    root = ET.fromstring(xml_bytes)
    paragraphs = []
    for paragraph in root.findall(".//w:p", _WORD_NAMESPACE):
        texts = [node.text for node in paragraph.findall(".//w:t", _WORD_NAMESPACE) if node.text]
        if texts:
            paragraphs.append("".join(texts))
    return "\n".join(paragraphs)


def extract_pptx_text(data: bytes) -> str:
    """Texto de cada slide, inclusive o das tabelas, numerado por slide."""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides_text = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.replace(" | ", "").strip():
                        parts.append(row_text)
        if parts:
            slides_text.append("--- Slide {} ---\n{}".format(number, "\n".join(parts)))
    return "\n\n".join(slides_text)


def extract_briefing_text(filename: str, data: bytes) -> Tuple[str, str]:
    """Devolve (texto_extraido, erro). Em caso de sucesso, erro é ""."""
    name = str(filename or "briefing")
    extension = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    if not data:
        return "", "O arquivo de briefing está vazio."

    try:
        if extension in {"txt", "md", "csv", "json"}:
            return decode_text_bytes(data).strip(), ""

        if extension == "docx":
            text = extract_docx_text(data).strip()
            if not text:
                return "", "Não foi possível extrair texto do .docx informado."
            return text, ""

        if extension == "pptx":
            text = extract_pptx_text(data).strip()
            if not text:
                return "", "Não foi possível extrair texto do .pptx informado."
            return text, ""

        return "", "Formato não suportado. Use .{}.".format(", .".join(BRIEFING_EXTENSIONS))
    except Exception as error:
        return "", "Erro ao processar briefing: {}".format(error)


def cap_text(text: str, max_chars: int = BRIEFING_MAX_CHARS) -> str:
    """Corta o texto guardado no banco, marcando que houve corte."""
    text = str(text or "").strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n...[briefing truncado no armazenamento]"
