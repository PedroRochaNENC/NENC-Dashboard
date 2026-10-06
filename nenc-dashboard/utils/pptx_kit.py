"""
Peças para apresentações PPTX dos módulos por projeto (Jornada de Compra, Teste Sensorial).

16:9, fundo branco, gráficos nativos do PowerPoint (botão direito → Editar
dados mexe nos números) e tabelas nativas. Cor por função: a entidade em
destaque em violeta NENC e as demais em cinza; texto sobre preenchimento
escolhe branco ou tinta pelo contraste. Valores ausentes ficam sem barra ou
como lacuna na linha, nunca viram zero.
"""

import io
import math
import re
from typing import Dict, List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from utils.pdf_report import numeric_column

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
LEFT = Inches(0.6)
CONTENT_WIDTH = Inches(12.13)
BODY_TOP = Inches(1.55)
BODY_BOTTOM = Inches(6.85)
FONT = "Calibri"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


INK = rgb("#1b1c24")
SECONDARY = rgb("#52514e")
MUTED = rgb("#898781")
RULE = rgb("#e1e0d9")
ACCENT = rgb("#5d5294")  # NENC accent-700: marca foco e títulos
FUNNEL = (rgb("#b5abfc"), rgb("#9184d9"), rgb("#5d5294"))  # accent-400 → accent → accent-700
BASE_BAR = rgb("#a3a6b3")
CATEGORICAL = (rgb("#2a78d6"), rgb("#eb6834"), rgb("#1baf7a"))  # azul, laranja, verde-água
OTHER = rgb("#a3a6b3")
HEADER_FILL = rgb("#eeedf6")
BAND_FILL = rgb("#f7f7f9")
WHITE = rgb("#ffffff")

MAX_TABLE_ROWS = 12
_CONTROL = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


# ---------------------------------------------------------------------------
# Valores
# ---------------------------------------------------------------------------

def clean_text(value) -> str:
    """Texto que o XML do Office aceita."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return _CONTROL.sub("", str(value))


def chart_value(value) -> Optional[float]:
    """NaN, infinito e vazio viram None: o gráfico deixa a barra em branco."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(number) or math.isinf(number) else number


def dash(text) -> str:
    return text if text not in ("", None) else "—"


def luminance(color: RGBColor) -> float:
    def channel(value: int) -> float:
        value = value / 255
        return value / 12.92 if value <= 0.03928 else ((value + 0.055) / 1.055) ** 2.4

    red, green, blue = color
    return 0.2126 * channel(red) + 0.7152 * channel(green) + 0.0722 * channel(blue)


def label_on(fill: RGBColor) -> RGBColor:
    """Branco ou tinta, o que tiver mais contraste com o preenchimento."""
    lum = luminance(fill)
    return WHITE if (1.05 / (lum + 0.05)) >= ((lum + 0.05) / (luminance(INK) + 0.05)) else INK


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------

def style_run(run, size: float, *, bold: bool = False, color: RGBColor = INK) -> None:
    font = run.font
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color


def add_runs(paragraph, text, size: float, *, color: RGBColor = INK, bold: bool = False) -> None:
    """Texto com **negrito** inline, em runs."""
    for index, part in enumerate(clean_text(text).split("**")):
        if part:
            run = paragraph.add_run()
            run.text = part
            style_run(run, size, bold=bold or index % 2 == 1, color=color)


def textbox(slide, left, top, width, height, text="", *, size: float = 14, bold: bool = False,
             color: RGBColor = INK, align=None, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    if text:
        paragraph = frame.paragraphs[0]
        if align is not None:
            paragraph.alignment = align
        add_runs(paragraph, text, size, color=color, bold=bold)
    return box


def bullet(paragraph, indent=Inches(0.28)) -> None:
    """Marcador com recuo pendente, na cor de acento."""
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(int(indent)))
    properties.set("indent", str(-int(indent)))
    color = OxmlElement("a:buClr")
    rgb = OxmlElement("a:srgbClr")
    rgb.set("val", str(ACCENT))
    color.append(rgb)
    properties.append(color)
    char = OxmlElement("a:buChar")
    char.set("char", "•")
    properties.append(char)


def write_items(frame, items: Sequence[Tuple[str, str]], size: float) -> None:
    """Itens ("h" título, "b" marcador, "p" parágrafo) numa caixa de texto."""
    for index, (kind, text) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if kind == "h":
            paragraph.space_before = Pt(0 if index == 0 else 10)
            paragraph.space_after = Pt(4)
            add_runs(paragraph, text, size + 2, color=ACCENT, bold=True)
        elif kind == "b":
            paragraph.space_after = Pt(6)
            bullet(paragraph)
            add_runs(paragraph, text, size)
        else:
            paragraph.space_after = Pt(8)
            add_runs(paragraph, text, size)


def markdown_items(text: str) -> List[Tuple[str, str]]:
    items = []
    in_table = False
    for raw in str(text or "").splitlines():
        line = raw.strip()
        if not line.startswith("|"):
            in_table = False
        if not line or line in ("---", "***", "___"):
            continue
        if line.startswith("|"):
            cells = [cell.strip().replace("**", "") for cell in line.strip("|").split("|")]
            if all(set(cell) <= set("-: ") for cell in cells):
                continue
            text_line = " · ".join(cells)
            items.append(("p", text_line if in_table else "**{}**".format(text_line)))
            in_table = True
        elif line.startswith("#"):
            items.append(("h", line.lstrip("#").strip().replace("**", "")))
        elif line[:2] in ("- ", "* ", "• "):
            items.append(("b", line[2:].strip()))
        else:
            items.append(("p", line))
    return items


# ---------------------------------------------------------------------------
# Deck
# ---------------------------------------------------------------------------

class Deck:
    def __init__(self, footer: str):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self.layout = self.prs.slide_layouts[6]  # em branco
        self.footer = footer

    def slide(self, title: str, subtitle: str = ""):
        slide = self.prs.slides.add_slide(self.layout)
        textbox(slide, LEFT, Inches(0.42), CONTENT_WIDTH, Inches(0.6), title, size=26, bold=True)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT, Inches(1.04), Inches(0.55), Inches(0.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        if subtitle:
            textbox(slide, LEFT, Inches(1.14), CONTENT_WIDTH, Inches(0.35), subtitle, size=12, color=SECONDARY)
        textbox(slide, LEFT, Inches(7.02), Inches(10.5), Inches(0.3), self.footer, size=9, color=MUTED)
        textbox(slide, Inches(11.73), Inches(7.02), Inches(1.0), Inches(0.3), str(len(self.prs.slides)),
                 size=9, color=MUTED, align=PP_ALIGN.RIGHT)
        return slide

    def flow(self, title: str, items: Sequence[Tuple[str, str]], *, subtitle: str = "",
             size: float = 16) -> None:
        """Itens de texto distribuídos em quantos slides precisarem.

        A conta é em linhas do corpo: a caixa tem 5,3 pol (381 pt) e cada linha
        ocupa 1,2 × o tamanho da fonte; cabem ~1.700/tamanho caracteres por
        linha. Uma linha fica de folga para a estimativa não estourar a caixa.
        """
        line_pt = size * 1.2
        capacity = (BODY_BOTTOM - BODY_TOP) / 12700 / line_pt - 1.0
        chars_per_line = int(1700 / size)
        spacing = {"h": 14 / line_pt, "b": 6 / line_pt, "p": 8 / line_pt}
        pages: List[List[Tuple[str, str]]] = []
        current: List[Tuple[str, str]] = []
        used = 0.0
        for kind, text in items:
            lines = max(1, math.ceil(len(clean_text(text)) / chars_per_line))
            cost = (lines * (size + 2) / size if kind == "h" else lines) + spacing.get(kind, 0.4)
            if current and used + cost > capacity:
                pages.append(current)
                current, used = [], 0.0
            current.append((kind, text))
            used += cost
        if current:
            pages.append(current)
        for index, page in enumerate(pages):
            slide = self.slide(title if index == 0 else "{} (continuação)".format(title),
                               subtitle if index == 0 else "")
            box = textbox(slide, LEFT, BODY_TOP, CONTENT_WIDTH, BODY_BOTTOM - BODY_TOP)
            write_items(box.text_frame, page, size)

    def table_slides(self, title: str, headers: Sequence[str], rows: Sequence[Sequence],
                     widths: Sequence[float], *, subtitle: str = "") -> None:
        for start in range(0, max(1, len(rows)), MAX_TABLE_ROWS):
            chunk = rows[start:start + MAX_TABLE_ROWS]
            if not chunk:
                return
            slide = self.slide(title if start == 0 else "{} (continuação)".format(title),
                               subtitle if start == 0 else "")
            native_table(slide, LEFT, BODY_TOP, headers, chunk, widths)

    def save(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


# ---------------------------------------------------------------------------
# Tabelas e números-chave
# ---------------------------------------------------------------------------

def table_cell(cell, value, size: float, *, bold: bool = False, fill: RGBColor = WHITE, align: str = "L") -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT if align == "R" else PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = clean_text(value)
    style_run(run, size, bold=bold, color=INK)


def native_table(slide, left, top, headers: Sequence[str], rows: Sequence[Sequence], widths: Sequence[float],
           *, size: float = 11, row_height: float = 0.34) -> None:
    """Tabela nativa: cabeçalho em tom de acento, linhas alternadas, números à direita."""
    align = ["R" if numeric_column([row[index] for row in rows]) else "L" for index in range(len(headers))]
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top,
                                   Inches(sum(widths)), Inches(row_height * (len(rows) + 1)))
    table = shape.table
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row in table.rows:
        row.height = Inches(row_height)
    for index, header in enumerate(headers):
        table_cell(table.cell(0, index), header, size, bold=True, fill=HEADER_FILL, align=align[index])
    for row_index, row in enumerate(rows, start=1):
        fill = WHITE if row_index % 2 else BAND_FILL
        for index, value in enumerate(row):
            table_cell(table.cell(row_index, index), value, size, fill=fill, align=align[index])


def kpis(slide, top, items: Sequence[Tuple[str, str]], *, height=Inches(1.25)) -> None:
    if not items:
        return
    gap = Inches(0.25)
    width = int((CONTENT_WIDTH - gap * (len(items) - 1)) / len(items))
    for index, (label, value) in enumerate(items):
        x = LEFT + index * (width + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RULE
        box.line.width = Pt(0.75)
        box.shadow.inherit = False
        textbox(slide, x + Inches(0.18), top + Inches(0.12), width - Inches(0.36), Inches(0.45), label,
                 size=12, color=SECONDARY)
        textbox(slide, x + Inches(0.18), top + Inches(0.6), width - Inches(0.36), Inches(0.55), value,
                 size=28, bold=True)


# ---------------------------------------------------------------------------
# Gráficos nativos
# ---------------------------------------------------------------------------

def chart_base(chart, title: str, *, legend: bool) -> None:
    chart.font.name = FONT
    chart.font.size = Pt(11)
    chart.font.color.rgb = SECONDARY
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
    if title:
        chart.has_title = True
        frame = chart.chart_title.text_frame
        frame.text = clean_text(title)
        for run in frame.paragraphs[0].runs:
            style_run(run, 13, bold=True)
    else:
        chart.has_title = False


def bar_axes(chart, maximum: Optional[float] = None) -> None:
    categories = chart.category_axis
    categories.reverse_order = True  # primeira categoria no topo, como na tela
    categories.has_major_gridlines = False
    categories.tick_labels.font.size = Pt(11)
    categories.format.line.color.rgb = RULE
    values = chart.value_axis
    values.has_major_gridlines = False
    values.visible = False  # os valores estão nos rótulos
    values.minimum_scale = 0
    if maximum:
        values.maximum_scale = maximum


def gap_width(height, categories: int, per_category: int = 1, *, legend: bool = False,
               bar_in: float = 0.3) -> int:
    """Espaço entre grupos (% da barra) para cada barra ter ~`bar_in` polegadas."""
    plot_in = height / 914400 - 0.45 - (0.4 if legend else 0.0)
    band = plot_in / max(1, categories)
    return int(max(40, min(400, (band / bar_in - per_category) * 100)))


def bar_chart(slide, box, categories, values, *, highlight=None, number_format: str = "0%",
               maximum: Optional[float] = None, title: str = "", color: RGBColor = ACCENT):
    x, y, width, height = box
    data = CategoryChartData(number_format=number_format)
    data.categories = [clean_text(c) for c in categories]
    data.add_series("Valor", [chart_value(v) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, width, height, data).chart
    chart_base(chart, title, legend=False)
    plot = chart.plots[0]
    plot.gap_width = gap_width(height, len(categories))
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(11)
    labels.font.color.rgb = SECONDARY
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color
    if highlight is not None:
        for index, flag in enumerate(highlight):
            fill = series.points[index].format.fill
            fill.solid()
            fill.fore_color.rgb = ACCENT if flag else BASE_BAR
    bar_axes(chart, maximum)
    return chart


def clustered_chart(slide, box, categories, series: Sequence[Tuple[str, Sequence]], colors, *,
                     number_format: str = "0%", maximum: Optional[float] = None, title: str = ""):
    x, y, width, height = box
    data = CategoryChartData(number_format=number_format)
    data.categories = [clean_text(c) for c in categories]
    for name, values in series:
        data.add_series(clean_text(name), [chart_value(v) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, width, height, data).chart
    chart_base(chart, title, legend=True)
    plot = chart.plots[0]
    plot.gap_width = gap_width(height, len(categories), len(series), legend=True, bar_in=0.22)
    plot.overlap = -8  # fresta entre as barras do grupo
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(10)
    labels.font.color.rgb = SECONDARY
    for chart_series, color in zip(plot.series, colors):
        chart_series.format.fill.solid()
        chart_series.format.fill.fore_color.rgb = color
    bar_axes(chart, maximum)
    return chart


def stacked_chart(slide, box, categories, series: Sequence[Tuple[str, Sequence]], colors, *, title: str = ""):
    """Barras 100% empilhadas; rótulo só em segmento que cabe (≥ 8%)."""
    x, y, width, height = box
    data = CategoryChartData(number_format="0%")
    data.categories = [clean_text(c) for c in categories]
    cleaned = [(name, [chart_value(v) for v in values]) for name, values in series]
    for name, values in cleaned:
        data.add_series(clean_text(name), values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, x, y, width, height, data).chart
    chart_base(chart, title, legend=True)
    plot = chart.plots[0]
    plot.gap_width = gap_width(height, len(categories), legend=True, bar_in=0.45)
    plot.overlap = 100
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0%"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.CENTER
    labels.font.size = Pt(11)
    labels.font.bold = True
    totals = [sum(values[i] or 0 for _, values in cleaned) for i in range(len(categories))]
    for (name, values), chart_series, color in zip(cleaned, plot.series, colors):
        chart_series.format.fill.solid()
        chart_series.format.fill.fore_color.rgb = color
        # Fresta branca entre os segmentos.
        chart_series.format.line.color.rgb = WHITE
        chart_series.format.line.width = Pt(1.5)
        series_labels = chart_series.data_labels
        series_labels.show_value = True
        series_labels.position = XL_LABEL_POSITION.CENTER
        series_labels.number_format = "0%"
        series_labels.number_format_is_linked = False
        series_labels.font.color.rgb = label_on(color)
        series_labels.font.size = Pt(11)
        series_labels.font.bold = True
        for index, value in enumerate(values):
            if value is None or not totals[index] or value / totals[index] < 0.08:
                chart_series.points[index].data_label.text_frame.text = ""
    bar_axes(chart, 1.0)
    return chart

def line_chart(slide, box, categories, series: Sequence[Tuple[str, Sequence]], colors, *,
               number_format: str = "0.00", title: str = "", minimum: Optional[float] = None,
               maximum: Optional[float] = None):
    """Linhas nativas (uma por série), para curvas no tempo; ausente vira lacuna, nunca zero."""
    x, y, width, height = box
    data = CategoryChartData(number_format=number_format)
    data.categories = [clean_text(c) for c in categories]
    for name, values in series:
        data.add_series(clean_text(name), [chart_value(v) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, x, y, width, height, data).chart
    chart_base(chart, title, legend=len(series) > 1)
    for chart_series, color in zip(chart.plots[0].series, colors):
        chart_series.smooth = False
        chart_series.marker.style = None
        line = chart_series.format.line
        line.color.rgb = color
        line.width = Pt(2)
    categories_axis = chart.category_axis
    categories_axis.has_major_gridlines = False
    categories_axis.tick_labels.font.size = Pt(10)
    categories_axis.format.line.color.rgb = RULE
    values_axis = chart.value_axis
    values_axis.has_major_gridlines = True
    values_axis.major_gridlines.format.line.color.rgb = RULE
    values_axis.tick_labels.font.size = Pt(10)
    values_axis.tick_labels.number_format = number_format
    values_axis.tick_labels.number_format_is_linked = False
    if minimum is not None:
        values_axis.minimum_scale = minimum
    if maximum is not None:
        values_axis.maximum_scale = maximum
    return chart


# ---------------------------------------------------------------------------
# Imagens
# ---------------------------------------------------------------------------

def picture(slide, content: bytes, box) -> None:
    """Imagem inteira dentro da caixa, centralizada, sem distorcer."""
    from PIL import Image

    left, top, width, height = box
    with Image.open(io.BytesIO(content)) as image:
        ratio = image.width / image.height
    if width / height > ratio:
        w, h = int(height * ratio), int(height)
    else:
        w, h = int(width), int(width / ratio)
    slide.shapes.add_picture(io.BytesIO(content), int(left + (width - w) / 2), int(top + (height - h) / 2), w, h)


def picture_slides(deck: Deck, pictures: Sequence[Dict], title: str, subtitle: str = "") -> None:
    """Duas imagens por slide, com a legenda embaixo."""
    gap = Inches(0.3)
    width = int((CONTENT_WIDTH - gap) / 2)
    height = BODY_BOTTOM - BODY_TOP - Inches(0.45)
    for start in range(0, len(pictures), 2):
        slide = deck.slide(title if start == 0 else "{} (continuação)".format(title), subtitle if start == 0 else "")
        for index, item in enumerate(pictures[start:start + 2]):
            left = LEFT + index * (width + gap)
            try:
                picture(slide, item["content"], (left, BODY_TOP, width, height))
            except Exception:
                continue
            textbox(slide, left, BODY_TOP + height + Inches(0.08), width, Inches(0.3), item["title"], size=11,
                     color=SECONDARY, align=PP_ALIGN.CENTER)
