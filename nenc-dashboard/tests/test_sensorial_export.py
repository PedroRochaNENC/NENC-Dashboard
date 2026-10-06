"""Exportações do Teste Sensorial: Excel com dicionário, PDF e PPTX sem nome de ninguém."""

import html
import io
import unittest
import zipfile

import openpyxl
import pandas as pd
from PIL import Image
from pptx import Presentation

from tests.test_pdf_report import _pdf_text
from tests.test_sensorial_ingest import CANARY_WORDS, _indicadores_csv
from tests.test_sensorial_model import BundleBase, _windows
from utils import sensorial_export, sensorial_ingest, sensorial_metrics, sensorial_quality
from utils.sensorial_pdf import build_pdf
from utils.sensorial_pptx import build_pptx

SESSIONS = ([("a{}".format(i), "P{:02d}".format(i), "2001A", "A") for i in range(1, 7)]
            + [("n{}".format(i), "P{:02d}".format(i), "2001Neutro", "N") for i in range(1, 7)]
            + [("b{}".format(i), "P{:02d}".format(i), "2001Basal", None) for i in range(1, 7)])
PERSONAL_COLUMNS = ("filename", "participante", "participante_original", "Participant Name", "computador")
ANALYSIS = {"id": 4, "mode": "rapida", "model": "m", "created_at": "2026-01-02 10:00", "data_version": 3,
            "analysis_text": "## Leitura\n\nA amostra A ficou acima do controle.", "search": [],
            "citations": [{"file_id": "file_1", "filename": "literatura.pdf", "quote": "Trecho.", "score": 0.8}]}


def _zip_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return html.unescape(" ".join(archive.read(name).decode("utf-8", "replace")
                                      for name in archive.namelist() if name.endswith(".xml")))


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (300, 300), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def _peripherals():
    rows = []
    for index, (sessao_id, code, experimento, amostra) in enumerate(SESSIONS):
        for stage in ("Basal", "Olfacao", "PosOlfacao"):
            rows.append({"sessao_id": sessao_id, "participant_code": code, "experimento": experimento,
                         "amostra": amostra, "Etapa": stage, "Bloco": 1, "Tempo": 0.0, "BPM": 70.0 + index,
                         "RMSSD": 30.0 + index, "GSR_CAL_mean": 2.0 + 0.1 * index, "BPM_zscore": 0.1 * index,
                         "RMSSD_zscore": -0.1 * index, "GSR_CAL_zscore": 0.2 * index, "qualidade_fc": "ok",
                         "qualidade_gsr": "ok"})
    return pd.DataFrame(rows)


def _trials():
    rows = []
    for sessao_id, code, experimento, amostra in SESSIONS[:6]:
        for trial, (word, rt, answer) in enumerate((("Fresco", 0.6, "Sim"), ("Leve", 1.4, "Sim"),
                                                    ("Doce", 1.0, "Nao"), ("Fresco", 0.8, "Sim"))):
            rows.append({"sessao_id": sessao_id, "participant_code": code, "experimento": experimento,
                         "amostra": amostra, "tentativa": trial + 1, "palavra": word, "rt": rt, "resposta": answer})
    return pd.DataFrame(rows)


class _Base(BundleBase):
    def setUp(self):
        super().setUp()
        self._file("eeg_psd", _windows(SESSIONS, stages=("Basal", "Olfacao", "PosOlfacao"), per_stage=6))
        # Uma tabela lida do pipeline com o nome-canário: ele não pode chegar a nenhum arquivo.
        self._file("eeg_indicadores", sensorial_ingest.parse_file("indicadores.csv", _indicadores_csv()).table)
        self._file("perifericos_metricas", _peripherals())
        self._file("associacao_tentativas", _trials())
        self.model = self._model(participants=[{"code": "P01", "profile": {"sexo": "F"},
                                                "notes": "=HYPERLINK(\"x\")"}])
        self.metrics = sensorial_metrics.compute_all(self.model)
        self.project = {"id": 9, "name": "Estudo Fragrância", "objetivo": "Comparar duas fragrâncias.",
                        "data_version": 3}
        self.quality = sensorial_quality.session_quality(self.model)

    def _files(self, model=None, metrics=None):
        model, metrics = model or self.model, metrics or self.metrics
        excel, excel_name = sensorial_export.build_excel(self.project, model, metrics, quality=self.quality,
                                                         analyses=[ANALYSIS], recorte="todo o projeto")
        pdf, pdf_name = build_pdf(self.project, model, metrics, recorte="todo o projeto", analysis=ANALYSIS,
                                  images=[(_png(), "2001A · Olfação")], generated_at="02/01/2026 10:00")
        pptx, pptx_name = build_pptx(self.project, model, metrics, recorte="todo o projeto", analysis=ANALYSIS,
                                     images=[{"content": _png(), "title": "2001A · Olfação"}])
        self.assertEqual([excel_name, pdf_name, pptx_name],
                         ["teste_sensorial_Estudo_Fragrancia_9.xlsx", "teste_sensorial_analise_Estudo_Fragrancia_9.pdf",
                          "teste_sensorial_analise_Estudo_Fragrancia_9.pptx"])
        return excel, pdf, pptx


class ExcelTests(_Base):
    def test_sheets_in_order_and_the_dictionary_describes_every_column(self):
        excel, _, _ = self._files()
        workbook = openpyxl.load_workbook(io.BytesIO(excel), read_only=True)
        self.assertEqual(workbook.sheetnames[:4], ["Projeto", "Condicoes", "Etapas", "Sessoes"])
        self.assertEqual(workbook.sheetnames[-1], "Dicionario")
        rows = list(workbook["Dicionario"].iter_rows(values_only=True))[1:]
        described = {(sheet, column): text for sheet, _, column, text in rows}
        for sheet in workbook.sheetnames[:-1]:
            header = next(workbook[sheet].iter_rows(values_only=True), ())
            for column in header:
                self.assertTrue(described.get((sheet, column)), "{}.{} sem descrição".format(sheet, column))

    def test_the_tables_relate_and_nothing_becomes_a_formula(self):
        excel, _, _ = self._files()
        workbook = openpyxl.load_workbook(io.BytesIO(excel), read_only=True)

        def sheet(name):
            rows = list(workbook[name].iter_rows(values_only=True))
            return pd.DataFrame(rows[1:], columns=rows[0])

        sessions, windows = sheet("Sessoes"), sheet("Janelas_EEG")
        self.assertTrue(set(windows["sessao_id"]) <= set(sessions["sessao_id"]))
        self.assertIn("qualidade", sessions.columns)
        self.assertEqual(sheet("Participantes").set_index("participant_code").loc["P01", "notas"],
                         "'=HYPERLINK(\"x\")")
        self.assertEqual(sheet("Analises_IA")["id"].tolist(), [4])
        self.assertEqual(sheet("Citacoes_IA")["documento"].tolist(), ["literatura.pdf"])
        self.assertIn("FAI", set(sheet("Comparacoes")["medida"]))
        self.assertEqual(set(sheet("Associacao")["palavra"]), {"Fresco", "Leve", "Doce"})


class PrivacyTests(_Base):
    def test_no_name_nor_personal_column_reaches_any_file(self):
        excel, pdf, pptx = self._files()
        texts = {"xlsx": _zip_text(excel).lower(), "pptx": _zip_text(pptx).lower(), "pdf": _pdf_text(pdf).lower()}
        for kind, text in texts.items():
            for word in CANARY_WORDS:
                self.assertNotIn(word, text, kind)
        workbook = openpyxl.load_workbook(io.BytesIO(excel), read_only=True)
        for name in workbook.sheetnames:
            header = next(workbook[name].iter_rows(values_only=True), ())
            personal = [c for c in header if c is not None and sensorial_ingest.is_personal_column(c)]
            self.assertEqual(personal, [], name)
            self.assertFalse(set(PERSONAL_COLUMNS) & set(header), name)


class ReportTests(_Base):
    def test_the_deck_has_native_charts_curves_pictures_and_the_analysis(self):
        _, _, pptx = self._files()
        deck = Presentation(io.BytesIO(pptx))
        self.assertEqual(round(deck.slide_width / deck.slide_height, 2), 1.78)
        titles = [next((s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text), "")
                  for slide in deck.slides]
        for title in ("Achados", "Valência emocional", "Teste de associação: Score por claim",
                      "Topomapas do pipeline", "Amostra", "Análise de IA", "Limitações"):
            self.assertIn(title, titles)
        charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        self.assertTrue(any(chart.chart_type is not None and "LINE" in str(chart.chart_type) for chart in charts))
        bars = [c for c in charts if "BAR" in str(c.chart_type)]
        self.assertTrue(bars)
        # Barra negativa (Score, assimetria) na cor da série, nunca em branco.
        self.assertTrue(all(not series.invert_if_negative for chart in bars for series in chart.plots[0].series))

    def test_the_pdf_carries_the_sections_and_the_chosen_analysis(self):
        _, pdf, _ = self._files()
        self.assertTrue(pdf.startswith(b"%PDF"))
        text = _pdf_text(pdf)
        for piece in ("TESTE SENSORIAL", "Estudo Fragr", "Achados", "Teste de associa", "Limita", "Leitura"):
            self.assertIn(piece, text)

    def test_tiny_scale_indices_go_in_scientific_notation(self):
        from utils import sensorial_pptx

        self.assertEqual(sensorial_pptx._number_format([("A", [1e-12, None, 3e-12])]), "0.00E+00")
        self.assertEqual(sensorial_pptx._number_format([("A", [0.4, -0.2])]), "0.00")

    def test_topomaps_of_the_analysed_stages_in_design_order(self):
        model = {"design": {"etapas": [{"codigo": "Basal", "rotulo": "Basal", "papel": "referencia"},
                                       {"codigo": "Olfacao", "rotulo": "Olfação", "papel": "exposicao"},
                                       {"codigo": "Fim", "rotulo": "Fim", "papel": "ignorar"}]},
                 "topomaps": [{"file_id": 1, "experimento": "2001B", "etapa_arquivo": "Olfacao"},
                              {"file_id": 2, "experimento": "2001A", "etapa_arquivo": "Olfacao"},
                              {"file_id": 3, "experimento": "2001A", "etapa_arquivo": "Basal"},
                              {"file_id": 4, "experimento": "2001A", "etapa_arquivo": "Fim"},
                              {"file_id": 5, "experimento": "2001A", "etapa_arquivo": "Audio_Gravacao"}]}
        self.assertEqual(sensorial_export.report_topomaps(model),
                         [{"file_id": 3, "title": "2001A · Basal"}, {"file_id": 2, "title": "2001A · Olfação"},
                          {"file_id": 1, "title": "2001B · Olfação"}])
        self.assertEqual(len(sensorial_export.report_topomaps(model, limit=1)), 1)


class EmptyProjectTests(_Base):
    def test_an_empty_project_still_gives_the_three_files(self):
        self.files = []
        model = self._model()
        metrics = sensorial_metrics.compute_all(model)
        self.quality = sensorial_quality.session_quality(model)
        excel, pdf, pptx = self._files(model, metrics)
        workbook = openpyxl.load_workbook(io.BytesIO(excel), read_only=True)
        self.assertIn("Dicionario", workbook.sheetnames)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreaterEqual(len(Presentation(io.BytesIO(pptx)).slides), 3)


if __name__ == "__main__":
    unittest.main()
