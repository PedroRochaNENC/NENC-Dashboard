import io
import unittest
import zipfile

from pptx import Presentation
from pptx.util import Inches

from utils.briefing import (
    cap_text,
    decode_text_bytes,
    extract_briefing_text,
    extract_pptx_text,
)


def _docx_bytes(*paragraphs: str) -> bytes:
    body = "".join(
        '<w:p><w:r><w:t>{}</w:t></w:r></w:p>'.format(text) for text in paragraphs
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>{}</w:body></w:document>".format(body)
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Objetivo do estudo"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Marca"
    table.cell(0, 1).text = "Meta"
    table.cell(1, 0).text = "Marca A"
    table.cell(1, 1).text = "Ser notada"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class BriefingExtractionTests(unittest.TestCase):
    def test_plain_text_in_any_common_encoding(self):
        self.assertEqual(decode_text_bytes("ação".encode("latin-1")), "ação")
        text, error = extract_briefing_text("contexto.md", "# Título\nação".encode("utf-8"))
        self.assertEqual((text, error), ("# Título\nação", ""))

    def test_docx_paragraphs_come_out_in_order(self):
        text, error = extract_briefing_text("b.docx", _docx_bytes("Primeiro", "Segundo"))
        self.assertEqual(error, "")
        self.assertEqual(text, "Primeiro\nSegundo")

    def test_pptx_keeps_slides_and_tables(self):
        text = extract_pptx_text(_pptx_bytes())
        self.assertIn("--- Slide 1 ---", text)
        self.assertIn("Objetivo do estudo", text)
        self.assertIn("Marca A | Ser notada", text)

    def test_unsupported_and_empty_files_explain_why(self):
        _, error = extract_briefing_text("planilha.xlsx", b"x")
        self.assertIn("Formato não suportado", error)
        _, error = extract_briefing_text("vazio.txt", b"")
        self.assertIn("vazio", error)

    def test_broken_docx_is_an_error_not_an_exception(self):
        text, error = extract_briefing_text("quebrado.docx", b"nao sou um zip")
        self.assertEqual(text, "")
        self.assertIn("Erro ao processar briefing", error)

    def test_cap_marks_the_cut(self):
        self.assertEqual(cap_text("  curto  "), "curto")
        capped = cap_text("x" * 30, max_chars=10)
        self.assertTrue(capped.startswith("x" * 10))
        self.assertIn("truncado", capped)


if __name__ == "__main__":
    unittest.main()
