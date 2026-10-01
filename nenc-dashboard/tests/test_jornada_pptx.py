"""Apresentação da Jornada: 16:9, gráficos nativos editáveis e ausentes sem virar zero."""

import io
import math
import unittest

import pandas as pd
from pptx import Presentation

from tests.test_jornada_model import SAMPLES, SECONDS, _bundle, _file, _frames
from utils import jornada_pptx
from utils.jornada_metrics import compute_all
from utils.jornada_model import build_model
from utils.jornada_quality import run_quality


def _titles(deck):
    titles = []
    for slide in deck.slides:
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
        titles.append(texts[0] if texts else "")
    return titles


class ValueTests(unittest.TestCase):
    def test_missing_values_become_blank_points(self):
        self.assertIsNone(jornada_pptx._number(float("nan")))
        self.assertIsNone(jornada_pptx._number(math.inf))
        self.assertIsNone(jornada_pptx._number(""))
        self.assertEqual(jornada_pptx._number("0.5"), 0.5)

    def test_labels_inside_fills_pick_the_readable_ink(self):
        self.assertEqual(jornada_pptx._label_on(jornada_pptx.ACCENT), jornada_pptx.WHITE)
        self.assertEqual(jornada_pptx._label_on(jornada_pptx.CATEGORICAL[1]), jornada_pptx.INK)

    def test_control_characters_are_dropped(self):
        self.assertEqual(jornada_pptx._clean("a\x00b\x1fc"), "abc")


class DeckTests(unittest.TestCase):
    def setUp(self):
        files = [
            _file(1, "DSP1234-INDIVIDUAL2.csv", SAMPLES),
            _file(2, "ATACADO-INDIVIDUAL2.csv", SECONDS),
            _file(3, "Pt01-JLivre-DSP1234.csv", _frames(200, 0.05)),
        ]
        self.bundle = _bundle(files)
        self.bundle["project"]["questions"] = "Visibilidade da marca foco"
        self.model = build_model(self.bundle)
        self.metrics = compute_all(self.model, {})

    def _deck(self, **kwargs):
        data, name = jornada_pptx.build_pptx(self.bundle["project"], self.model, self.metrics,
                                             generated_at="01/01/2026 10:00", **kwargs)
        self.assertTrue(name.endswith(".pptx"))
        return Presentation(io.BytesIO(data))

    def test_the_deck_is_wide_and_carries_native_charts(self):
        deck = self._deck(quality=run_quality(self.model, self.bundle["project"]))
        self.assertEqual((deck.slide_width, deck.slide_height), (jornada_pptx.SLIDE_WIDTH, jornada_pptx.SLIDE_HEIGHT))
        titles = _titles(deck)
        for expected in ("Números-chave", "Principais achados", "Share visual por marca", "Amostra e qualidade",
                         "Limitações"):
            self.assertIn(expected, titles)
        charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        self.assertTrue(charts)
        share = charts[0]
        # Os dados do gráfico continuam editáveis e batem com a tabela de métricas.
        brand = self.metrics["brand"]
        first_cell = brand[brand["cell"] == brand["cell"].iloc[0]].sort_values("share_mean", ascending=False)
        values = share.plots[0].series[0].values
        for got, expected in zip(values, first_cell["share_mean"]):
            self.assertAlmostEqual(got, expected)

    def test_choices_and_pictures_get_their_slides(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        from tests.test_jornada_choice import _field_log
        from tests.test_jornada_import_review import _png

        bundle = dict(self.bundle, files=self.bundle["files"] + [_file(9, "Relação Coletas.xlsx", _field_log())])
        bundle["project"] = dict(bundle["project"], marcas="Alfa\nBeta\nGama Livre")
        model = build_model(bundle)
        pictures = [{"group": "loja", "title": "Heatmap · DSP 1234", "content": _png()},
                    {"group": "marca", "title": "Embalagem · Alfa", "content": _png()}]
        data, _ = jornada_pptx.build_pptx(bundle["project"], model, compute_all(model, {}), images=pictures)
        deck = Presentation(io.BytesIO(data))
        titles = _titles(deck)
        for expected in ("Escolha — Jornada Estimulada", "Da atenção à escolha", "A gôndola de cada loja",
                         "As embalagens"):
            self.assertIn(expected, titles)
        placed = [shape for slide in deck.slides for shape in slide.shapes
                  if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(placed), 2)
        # Sem distorção: a proporção da imagem (640 × 480) se mantém.
        self.assertAlmostEqual(placed[0].width / placed[0].height, 640 / 480, places=2)

    def test_a_missing_value_leaves_the_bar_blank(self):
        slide = jornada_pptx._Deck("rodapé").slide("Teste")
        chart = jornada_pptx._bar_chart(slide, (0, 0, jornada_pptx.Inches(4), jornada_pptx.Inches(3)),
                                        ["A", "B"], [0.4, float("nan")])
        self.assertEqual(chart.plots[0].series[0].values, (0.4, None))

    def test_long_text_continues_on_the_next_slide(self):
        deck = jornada_pptx._Deck("rodapé")
        deck.flow("Achados", [("b", "Frase de teste número {} com algum texto.".format(i)) for i in range(60)])
        titles = _titles(deck.prs)
        self.assertGreater(len(titles), 1)
        self.assertEqual(titles[1], "Achados (continuação)")

    def test_the_analysis_goes_in_with_its_references(self):
        deck = self._deck(analysis={
            "model": "m", "mode": "rapida", "data_version": 3,
            "analysis_text": "## Síntese\n\nMarca A **lidera**.\n\n| Loja | Share |\n|---|---|\n| 1 | 40% |",
            "citations": [{"filename": "briefing.docx", "quote": "trecho"}],
        })
        texts = "\n".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes
                          if shape.has_text_frame)
        self.assertIn("Análise de IA", texts)
        self.assertIn("Loja · Share", texts)
        self.assertIn("briefing.docx", texts)

    def test_an_empty_project_still_gives_a_deck(self):
        model = build_model(_bundle([]))
        data, _ = jornada_pptx.build_pptx({"id": 5, "name": "Vazio"}, model, compute_all(model, {}))
        deck = Presentation(io.BytesIO(data))
        self.assertIn("Principais achados", _titles(deck))

    def test_attributes_follow_the_configured_order(self):
        metrics = dict(self.metrics)
        metrics["attributes"] = pd.DataFrame([
            {"cell": "Loja 1", "dimension": "tipo", "value": "Noturno", "share_mean": 0.7, "reach": 1.0,
             "n": 3, "n_defined": 3, "task": "livre", "store": "1", "store_label": "Loja 1"},
            {"cell": "Loja 1", "dimension": "tipo", "value": "Diurno", "share_mean": 0.3, "reach": 1.0,
             "n": 3, "n_defined": 3, "task": "livre", "store": "1", "store_label": "Loja 1"},
        ])
        data, _ = jornada_pptx.build_pptx(self.bundle["project"], self.model, metrics)
        deck = Presentation(io.BytesIO(data))
        stacked = [shape.chart for slide in deck.slides for shape in slide.shapes
                   if shape.has_chart and len(shape.chart.plots[0].series) == 2
                   and shape.chart.plots[0].series[0].name in ("Diurno", "Noturno")]
        self.assertEqual([s.name for s in stacked[0].plots[0].series], ["Diurno", "Noturno"])


if __name__ == "__main__":
    unittest.main()
