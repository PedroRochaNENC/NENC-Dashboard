"""Exportações da Jornada: o Excel reabre, as chaves relacionam e nada vira fórmula."""

import io
import math
import unittest

import openpyxl
import pandas as pd

from tests.test_jornada_model import SAMPLES, SECONDS, _bundle, _file, _frames
from utils import jornada_export
from utils.jornada_metrics import compute_all
from utils.jornada_model import build_model
from utils.jornada_quality import run_quality


def _sheet(workbook, name):
    rows = list(workbook[name].iter_rows(values_only=True))
    return pd.DataFrame(rows[1:], columns=rows[0])


class ExcelTests(unittest.TestCase):
    def setUp(self):
        files = [
            _file(1, "DSP1234-INDIVIDUAL2.csv", SAMPLES),
            _file(2, "ATACADO-INDIVIDUAL2.csv", SECONDS),
            _file(3, "Pt01-JLivre-DSP1234.csv", _frames(200, 0.05)),
            _file(4, "Pt03-JEstimulada-ATACADO.csv", _frames(601, 0.05)),
        ]
        self.bundle = _bundle(files, participants=[
            {"code": "Pt01", "profile": "P1", "tempo_informado": "1m21s", "notes": "=HYPERLINK(\"x\")"},
        ])
        self.model = build_model(self.bundle)
        self.metrics = compute_all(self.model, {"tasks": ["livre"]})
        self.quality = run_quality(self.model, self.bundle["project"])

    def _workbook(self, **kwargs):
        data, name = jornada_export.build_excel(
            self.bundle["project"], self.model, self.metrics, quality=self.quality, **kwargs)
        self.assertTrue(name.endswith(".xlsx"))
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True)

    def test_sheets_open_in_order_and_keys_relate(self):
        workbook = self._workbook(media=[{"participant_code": "Pt01", "task": "livre", "store": "1234"}])
        self.assertEqual(workbook.sheetnames[:4], ["Projeto", "Lojas", "Participantes", "Gravacoes"])
        self.assertEqual(workbook.sheetnames[-1], "Dicionario")
        recordings = _sheet(workbook, "Gravacoes")
        gaze = _sheet(workbook, "Olhar_AOI")
        catalog = _sheet(workbook, "Catalogo_AOI")
        self.assertTrue(set(gaze["recording_key"]) <= set(recordings["recording_key"]))
        self.assertTrue(set(gaze["aoi_key"]) <= set(catalog["aoi_key"]))
        # Olhar convertido para segundos: 40 amostras a 0,05 s.
        row = gaze[(gaze["participant"] == "Pt01") & (gaze["aoi"] == "Marca A Noturno p1")].iloc[0]
        self.assertAlmostEqual(row["dwell_s"], 2.0, places=3)
        video = recordings.set_index("recording_key")["has_video"]
        self.assertTrue(video["Pt01|livre|1234"])
        self.assertFalse(video["Pt03|estimulada|atacado"])
        self.assertIn("quality_label", recordings.columns)

    def test_the_project_sheet_records_the_filter(self):
        project = _sheet(self._workbook(), "Projeto")
        self.assertIn("Tarefas: Jornada Livre", project.loc[0, "recorte"])
        # Metricas seguem o recorte; as tabelas de dados vao completas.
        workbook = self._workbook()
        self.assertEqual(set(_sheet(workbook, "Metricas_Marca")["task"]), {"livre"})
        self.assertIn("estimulada", set(_sheet(workbook, "Gravacoes")["task"]))

    def test_text_never_becomes_a_formula_and_long_text_is_cut(self):
        long_text = "a" * 40_000
        workbook = self._workbook(analyses=[{
            "id": 1, "model": "m", "mode": "rapida", "data_version": 3, "analysis_text": long_text,
            "citations": [{"file_id": "f", "filename": "=x.docx", "quote": "trecho"}, "solta"],
        }])
        participants = _sheet(workbook, "Participantes").set_index("participant")
        self.assertTrue(participants.loc["Pt01", "notes"].startswith("'="))
        analyses = _sheet(workbook, "Analises_IA")
        self.assertLessEqual(len(analyses.loc[0, "analysis_text"]), 32_767)
        self.assertTrue(analyses.loc[0, "is_current"])
        citations = _sheet(workbook, "Citacoes")
        self.assertEqual(citations["filename"].tolist()[0], "'=x.docx")
        self.assertEqual(citations["quote"].tolist()[1], "solta")

    def test_infinity_is_written_as_blank(self):
        self.metrics["decision"] = pd.DataFrame([{"group_type": "loja", "group": "X", "n": 1,
                                                  "median_s": math.inf, "q1_s": 1.0, "q3_s": 1.0,
                                                  "min_s": 1.0, "max_s": 1.0}])
        decision = _sheet(self._workbook(), "Tempo_Decisao")
        self.assertIsNone(decision.loc[0, "median_s"])

    def test_every_column_is_in_the_dictionary(self):
        tables = jornada_export.excel_tables(self.bundle["project"], self.model, self.metrics,
                                             quality=self.quality)
        dictionary = tables["Dicionario"]
        described = set(zip(dictionary["aba"], dictionary["coluna"]))
        for sheet, frame in tables.items():
            if sheet == "Dicionario":
                continue
            for column in frame.columns:
                self.assertIn((sheet, column), described)
        blank = dictionary[dictionary["descricao"] == ""]
        self.assertTrue(blank.empty, blank[["aba", "coluna"]].values.tolist())

    def test_choices_and_times_go_as_plain_text_and_numbers(self):
        from tests.test_jornada_choice import _field_log

        bundle = dict(self.bundle, files=self.bundle["files"] + [_file(9, "Relação Coletas.xlsx", _field_log())])
        bundle["project"] = dict(bundle["project"], marcas="Alfa\nBeta\nGama Livre")
        model = build_model(bundle)
        metrics = compute_all(model, {})
        tables = jornada_export.excel_tables(bundle["project"], model, metrics, quality=self.quality)
        choices = tables["Escolhas"].set_index(["participant", "task"])
        self.assertEqual(choices.loc[("Pt03", "estimulada"), "chosen_brands"], "Gama Livre")
        self.assertEqual(choices.loc[("Pt03", "estimulada"), "considered_brands"], "Alfa, Gama Livre")
        self.assertNotIn("chosen_values", tables["Escolhas"].columns)
        self.assertIn("campo", set(tables["Tempos"]["source"]))
        self.assertFalse(tables["Escolha_Marca"].empty)
        workbook = openpyxl.load_workbook(io.BytesIO(jornada_export.write_workbook(tables, 1000)), read_only=True)
        self.assertIn("Escolhas", workbook.sheetnames)
        dictionary = tables["Dicionario"]
        blank = dictionary[dictionary["descricao"] == ""]
        self.assertTrue(blank.empty, blank[["aba", "coluna"]].values.tolist())

    def test_an_empty_project_still_has_every_header(self):
        model = build_model(_bundle([]))
        metrics = compute_all(model, {})
        tables = jornada_export.excel_tables({"id": 9, "name": "Vazio"}, model, metrics,
                                             quality=run_quality(model))
        for sheet in ("Gravacoes", "Olhar_AOI_bruto", "Comparacoes", "Elementos_Embalagem",
                      "Qualidade", "Achados"):
            self.assertTrue(len(tables[sheet].columns), sheet)
        data, _ = jornada_export.build_excel({"id": 9, "name": "Vazio"}, model, metrics)
        self.assertEqual(data[:2], b"PK")


class FilterTextTests(unittest.TestCase):
    def test_describes_the_filter_in_one_line(self):
        model = {"stores": pd.DataFrame({"store": ["2250"], "label": ["DSP 2250"]})}
        text = jornada_export.filters_text(
            {"tasks": ["livre"], "stores": ["2250"], "kinds": ["produto", "preco"]}, model)
        self.assertEqual(
            text, "Tarefas: Jornada Livre · Lojas: DSP 2250 · Perfis: todos · Preço: parte da atenção da marca")
        self.assertIn("Tarefas: todas", jornada_export.filters_text({}))


if __name__ == "__main__":
    unittest.main()
