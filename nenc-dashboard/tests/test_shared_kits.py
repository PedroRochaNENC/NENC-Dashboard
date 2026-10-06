"""Peças compartilhadas pelos módulos por projeto: planilha, cores e PPTX."""

import io
import math
import unittest

import pandas as pd
from pptx import Presentation

from utils import excel_export, pptx_kit
from utils.chart_kit import BRAND_SEQUENCE, OTHER_COLOR, color_map


class ExcelPiecesTests(unittest.TestCase):
    def test_clean_frame_keeps_headers_drops_infinity_and_caps_text(self):
        empty = excel_export.clean_frame(None, ["a", "b"])
        self.assertEqual(list(empty.columns), ["a", "b"])
        frame = excel_export.clean_frame(pd.DataFrame({"v": [1.0, math.inf], "t": ["x", "y" * 40_000]}))
        self.assertTrue(math.isnan(frame.loc[1, "v"]))
        self.assertLessEqual(len(frame.loc[1, "t"]), excel_export.EXCEL_CELL_MAX_CHARS)
        self.assertTrue(frame.loc[1, "t"].endswith("Excel]"))

    def test_analysis_rows_mark_the_current_version_and_read_any_citation(self):
        analyses, citations = excel_export.analysis_rows([
            {"id": 1, "data_version": 3, "citations": [{"filename": "a.pdf", "quote": "q"}, "solta"]},
            {"id": 2, "data_version": 2},
        ], data_version=3)
        self.assertEqual([row["is_current"] for row in analyses], [True, False])
        self.assertEqual([(c["filename"], c["quote"]) for c in citations], [("a.pdf", "q"), (None, "solta")])

    def test_dictionary_uses_descriptions_then_the_fallback(self):
        tables = {"Aba": pd.DataFrame(columns=["conhecida", "attr_x", "outra"])}
        frame = excel_export.dictionary_frame(
            tables, {"Aba": "Uma aba", "Dicionario": "Esta aba"}, {"conhecida": "Descrita"},
            lambda column: "Atributo" if column.startswith("attr_") else "")
        self.assertEqual(frame["descricao"].tolist(), ["Descrita", "Atributo", "", "Aba, coluna e o que ela contém."])
        self.assertEqual(frame["descricao_aba"].iloc[-1], "Esta aba")


class ColorTests(unittest.TestCase):
    def test_the_focus_takes_the_first_slot_and_the_ninth_goes_grey(self):
        names = ["n{}".format(i) for i in range(9)]
        colors = color_map(names, focus="N3")
        self.assertEqual(colors["n3"], BRAND_SEQUENCE[0])
        self.assertEqual(colors["n0"], BRAND_SEQUENCE[1])
        self.assertEqual(colors["n8"], OTHER_COLOR)
        # A cor segue a entidade: tirar uma não repinta as outras na mesma posição relativa.
        self.assertEqual(color_map(["a", "b"])["b"], color_map(["a", "b", "c"])["b"])


class DeckTests(unittest.TestCase):
    def test_a_deck_with_table_bars_line_and_picture_opens(self):
        from PIL import Image

        buffer = io.BytesIO()
        Image.new("RGB", (400, 200), (10, 20, 30)).save(buffer, format="PNG")
        deck = pptx_kit.Deck("rodapé")
        slide = deck.slide("Título", "Subtítulo")
        pptx_kit.native_table(slide, pptx_kit.LEFT, pptx_kit.BODY_TOP, ["A", "B"], [["x", "1"]], [2.0, 1.0])
        pptx_kit.bar_chart(slide, (pptx_kit.LEFT, pptx_kit.BODY_TOP, pptx_kit.Inches(5), pptx_kit.Inches(3)),
                           ["a", "b"], [0.5, None], highlight=[True, False])
        line = pptx_kit.line_chart(
            slide, (pptx_kit.LEFT, pptx_kit.BODY_TOP, pptx_kit.Inches(6), pptx_kit.Inches(3)),
            ["0,0", "0,5", "1,0"], [("A", [1.0, None, 2.0]), ("B", [0.5, 0.7, 0.9])],
            [pptx_kit.ACCENT, pptx_kit.BASE_BAR])
        self.assertEqual(list(line.plots[0].series[0].values), [1.0, None, 2.0])  # lacuna, não zero
        pptx_kit.picture_slides(deck, [{"content": buffer.getvalue(), "title": "Imagem"}], "Imagens")
        opened = Presentation(io.BytesIO(deck.save()))
        self.assertEqual(len(opened.slides), 2)
        picture = [s for s in opened.slides[1].shapes if s.shape_type == 13][0]
        self.assertAlmostEqual(picture.width / picture.height, 2.0, places=2)


if __name__ == "__main__":
    unittest.main()
